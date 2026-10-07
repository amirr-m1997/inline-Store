"""نمایش‌های سفارشی پنل: گزارش‌ساز با خروجی اکسل، سفارش سریع، اعلان‌ها و API خلاصه."""
from __future__ import annotations

import re
from datetime import timedelta
from io import BytesIO

from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.db import models
from django.shortcuts import redirect, render
from django.urls import reverse
from urllib.parse import urlencode
from django.utils import timezone

from customers.models import Company
from inventory.services import stock_check
from orders.models import Order, OrderLine
from orders.services import add_event, evaluate_approvals
from pricing.services import resolve_price

from . import analytics
from .charts import bars, hbars
from .models import Notification
from .utils import en, jalali, money_short, num, parse_jalali

DIMENSIONS = [
    ("province", "استان"),
    ("category", "دسته کالا"),
    ("brand", "برند"),
    ("product", "محصول"),
    ("company", "مشتری"),
    ("rep", "کارشناس فروش"),
    ("payment", "روش پرداخت"),
    ("month", "ماه"),
]


# --------------------------------------------------------------- گزارش‌ساز


def _report_filters(request) -> dict:
    today = timezone.localdate()
    start = parse_jalali(request.GET.get("from")) or (today - timedelta(days=30))
    end = parse_jalali(request.GET.get("to")) or today
    dimension = request.GET.get("dimension") or "province"
    if dimension not in dict(DIMENSIONS):
        dimension = "province"
    company_id = request.GET.get("company") or ""
    status = request.GET.get("status") or ""
    return {
        "start": start, "end": end, "dimension": dimension,
        "company_id": company_id, "status": status,
        "company": Company.objects.filter(pk=company_id).first() if company_id.isdigit() else None,
    }


def report_view(request):
    filters = _report_filters(request)
    report = analytics.sales_by(
        dimension=filters["dimension"], start=filters["start"], end=filters["end"],
        company=filters["company"], status=filters["status"] or None,
    )
    summary = analytics.report_summary(filters["start"], filters["end"])
    chart = bars([
        {"label": row["label"][:14], "value": row["amount"] / 1_000_000, "label_value": money_short(row["amount"])}
        for row in report["rows"][:8]
    ])
    daily = analytics.sales_series(max((filters["end"] - filters["start"]).days + 1, 1))

    context = {
        **__import__("django.contrib.admin", fromlist=["site"]).site.each_context(request),
        "title": "گزارش‌ها و خروجی اکسل",
        "filters": filters,
        "dimensions": DIMENSIONS,
        "report": report,
        "summary": summary,
        "chart": chart,
        "companies": Company.objects.order_by("name")[:300],
        "status_choices": Order._meta.get_field("status").choices,
        "from_jalali": jalali(filters["start"]),
        "to_jalali": jalali(filters["end"]),
        "query_string": request.GET.urlencode(),
    }
    return render(request, "core/reports.html", context)


def report_export_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    filters = _report_filters(request)
    report = analytics.sales_by(
        dimension=filters["dimension"], start=filters["start"], end=filters["end"],
        company=filters["company"], status=filters["status"] or None,
    )
    summary = analytics.report_summary(filters["start"], filters["end"])

    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش فروش"
    ws.sheet_view.rightToLeft = True
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0E7490")

    ws.append([f"گزارش فروش — {report['title']}"])
    ws.append([f"از {jalali(filters['start'])} تا {jalali(filters['end'])}"])
    ws.append([])
    ws.append(["ردیف", report["title"], "تعداد سفارش", "مبلغ (تومان)", "سهم درصد"])
    for cell in ws[4]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for index, row in enumerate(report["rows"], start=1):
        ws.append([index, row["label"], row["orders"], row["amount"], row["share_pct"]])
    ws.append([])
    ws.append(["جمع کل", "", sum(r["orders"] for r in report["rows"]), report["total"], 100])
    ws.append([])
    ws.append(["خلاصه دوره"])
    for label, value in [
        ("فروش کل (تومان)", summary["sales"]),
        ("تعداد سفارش", summary["orders"]),
        ("میانگین سبد", summary["avg_basket"]),
        ("بهای تمام‌شده", summary["cost"]),
        ("سود ناخالص", summary["gross_profit"]),
        ("حاشیه سود (٪)", summary["margin_pct"]),
        ("تخفیف اعطاشده", summary["discount"]),
        ("مشتریان فعال", summary["active_customers"]),
        ("مبلغ فاکتورشده", summary["invoiced"]),
        ("مبلغ وصول‌شده", summary["collected"]),
    ]:
        ws.append([label, value])

    widths = [8, 40, 14, 22, 12]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"sales-report-{filters['start']}-{filters['end']}.xlsx"
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# --------------------------------------------------------------- سفارش سریع

LINE_RE = re.compile(r"^\s*(?P<code>[A-Za-z0-9\-_/\.]+)\s*[×xX,\t ]\s*(?P<qty>[\d۰-۹.,]+)\s*$")


def _parse_quick_order(text: str) -> list[tuple[str, float]]:
    rows = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        match = LINE_RE.match(line)
        if not match:
            rows.append((line, 0))
            continue
        qty_text = en(match.group("qty")).replace(",", "")
        try:
            qty = float(qty_text)
        except ValueError:
            qty = 0
        rows.append((match.group("code").strip(), qty))
    return rows


def quick_order_view(request):
    companies = Company.objects.filter(is_active=True).order_by("name")
    selected = None
    company_id = request.POST.get("company") or request.GET.get("company") or ""
    if company_id.isdigit():
        selected = companies.filter(pk=company_id).first()

    text = request.POST.get("lines", "")
    po_number = request.POST.get("po_number", "")
    action = request.POST.get("action", "")
    rows: list[dict] = []
    matched_products = []

    if request.method == "POST" and text.strip():
        from catalog.models import Product

        for code, qty in _parse_quick_order(text):
            display_code = code if qty else ""
            product = (
                Product.objects.filter(code__iexact=code).first()
                or Product.objects.filter(name__iexact=code).first()
                if code
                else None
            )
            row = {"code": code or "—", "qty": qty, "product": product, "status": "", "key": "muted",
                   "note": "", "unit_price": 0, "list_price": 0, "total": 0}
            if product is None:
                row.update(status="کد ناشناس", key="danger", note="کالایی با این کد یافت نشد.")
                rows.append(row)
                continue
            row["product_name"] = product.name
            row["uom"] = product.get_uom_display()
            if qty <= 0:
                row.update(status="تعداد نامعتبر", key="danger")
                rows.append(row)
                continue
            notes: list[str] = []
            packed = product.apply_packaging(qty)
            if packed != qty:
                notes.append(f"مضرب بسته‌بندی: {num(qty, 0)} ← {num(packed, 0)}")
                qty = packed
                row["qty"] = packed
            error = product.validate_qty(qty)
            if error:
                row.update(status="خارج از قواعد فروش", key="warn", note="؛ ".join(notes + [error]))
                rows.append(row)
                continue
            price = resolve_price(product, company=selected, qty=qty)
            stock = stock_check(product, qty)
            row.update(
                list_price=price.list_price, unit_price=price.unit_price,
                total=int(price.unit_price * qty), source=price.source,
                discount_pct=price.total_discount_pct,
            )
            if not stock["sufficient"]:
                notes.append(f"موجودی آزاد {num(stock['free'], 0)} — کسری "
                             f"{num(stock['shortage'], 0)} (تأمین {num(product.lead_time_days)} روز)")
                row.update(status="کمبود موجودی", key="warn")
            else:
                row.update(status="یافته شد", key="ok")
            row["note"] = "؛ ".join(notes)
            matched_products.append((product, qty, price.unit_price, stock))
            rows.append(row)

        if action == "create" and selected and matched_products:
            order = Order.objects.create(
                company=selected, po_number=po_number, sales_rep=selected.sales_rep,
                created_by=request.user, payment_method="cheque" if selected.payment_terms == "cheque" else "credit",
                payment_terms=selected.get_payment_terms_display(),
                placed_by_rep_for_company=True,
                customer_note="ثبت‌شده از صفحه سفارش سریع پنل",
            )
            for product, qty, unit_price, stock in matched_products:
                OrderLine.objects.create(
                    order=order, product=product, qty=qty, unit_price=unit_price,
                    list_price=product.base_price, warehouse=stock["warehouse"],
                    lead_time_days=product.lead_time_days,
                )
            evaluate_approvals(order, actor=request.user)
            add_event(order, "ثبت از سفارش سریع", f"{len(matched_products)} قلم", actor=request.user)
            messages.success(request, f"سفارش {order.number} ساخته شد و در گردش تأیید قرار گرفت.")
            return redirect(reverse("admin:orders_order_change", args=[order.pk]))

        if action == "create" and not matched_products:
            messages.error(request, "هیچ ردیف قابل ثبت‌ای وجود ندارد؛ ابتدا کدها را اصلاح کنید.")

    context = {
        **__import__("django.contrib.admin", fromlist=["site"]).site.each_context(request),
        "title": "سفارش سریع (Quick Order)",
        "companies": companies,
        "selected": selected,
        "text": text,
        "po_number": po_number,
        "rows": rows,
        "matched": len(matched_products),
        "total_amount": sum(row.get("total", 0) for row in rows),
    }
    return render(request, "core/quick_order.html", context)


# --------------------------------------------------------------- اعلان‌ها


def alerts_view(request):
    """مرکز اعلان‌ها: فیلتر بر اساس شدت/نوع/وضعیت + جست‌وجو + عملیات گروهی."""
    if request.method == "POST":
        ids = request.POST.getlist("ids")
        action = request.POST.get("action", "mark_read")
        queryset = Notification.objects.filter(pk__in=ids) if ids else Notification.objects.all()
        if action == "mark_unread":
            count = queryset.update(is_read=False)
            messages.success(request, f"{num(count)} اعلان به «خوانده‌نشده» برگشت.")
        else:
            count = queryset.update(is_read=True)
            messages.success(request, f"{num(count)} اعلان خوانده‌شده علامت خورد.")
        return redirect(request.POST.get("next") or reverse("panel-alerts"))

    level = request.GET.get("level", "")
    kind = request.GET.get("kind", "")
    state = request.GET.get("state", "unread")
    query = request.GET.get("q", "").strip()

    queryset = Notification.objects.select_related("company")
    if level:
        queryset = queryset.filter(level=level)
    if kind:
        queryset = queryset.filter(kind=kind)
    if state == "unread":
        queryset = queryset.filter(is_read=False)
    elif state == "read":
        queryset = queryset.filter(is_read=True)
    if query:
        queryset = queryset.filter(
            models.Q(title__icontains=query) | models.Q(body__icontains=query)
        )

    base = Notification.objects.all()
    level_counts = {
        key: base.filter(level=key, is_read=False).count()
        for key, _label in Notification.Level.choices
    }
    context = {
        **__import__("django.contrib.admin", fromlist=["site"]).site.each_context(request),
        "title": "اعلان‌ها و هشدارها",
        "notifications": queryset.order_by("-created_at")[:200],
        "unread_count": base.filter(is_read=False).count(),
        "total_count": base.count(),
        "levels": Notification.Level.choices,
        "kinds": Notification.Kind.choices,
        "level_counts": level_counts,
        "filters": {"level": level, "kind": kind, "state": state, "q": query},
        "query_string": urlencode({k: v for k, v in
                                   {"level": level, "kind": kind, "state": state, "q": query}.items()
                                   if v}),
    }
    return render(request, "core/alerts.html", context)


# --------------------------------------------------------------- API خلاصه


def api_summary(request):
    """خلاصه‌ی KPI برای اتصال به داشبوردهای بیرونی/CRM."""
    kpi = analytics.kpis()
    return JsonResponse(
        {
            "generated_at": timezone.now().isoformat(),
            "kpi": {key: (int(value) if isinstance(value, (int, float)) else value)
                    for key, value in kpi.items()},
            "monthly_sales": analytics.monthly_sales(6),
            "top_products": analytics.top_products(5),
        },
        json_dumps_params={"ensure_ascii": False},
    )
