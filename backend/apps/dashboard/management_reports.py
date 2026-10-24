"""Database-backed management reporting for the internal admin console."""

from collections import defaultdict
from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Max, Q
from django.utils import timezone

from apps.catalog.models import Product
from apps.inventory.models import Receipt
from apps.orders.models import Order, OrderItem


COMPLETED_SALE_STATUSES = (
    Order.Status.CONFIRMED,
    Order.Status.PREPARING,
    Order.Status.SHIPPED,
    Order.Status.DELIVERED,
)
ZERO = Decimal("0.00")


def parse_report_range(params):
    today = timezone.localdate()
    default_start = today - timedelta(days=29)

    def parse(value, fallback):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date() if value else fallback
        except (TypeError, ValueError):
            return fallback

    start = parse(params.get("start"), default_start)
    end = parse(params.get("end"), today)
    if start > end:
        start, end = end, start
    stale_days = params.get("stale_days", "90")
    try:
        stale_days = max(30, min(int(stale_days), 730))
    except (TypeError, ValueError):
        stale_days = 90
    return start, end, stale_days


def _customer_name(order):
    if order.customer_company_name.strip():
        return order.customer_company_name.strip()
    if order.customer:
        company_name = (order.customer.company_name or "").strip()
        if company_name:
            return company_name
        full_name = order.customer.get_full_name().strip()
        return full_name or order.customer.email or order.customer.phone or "مشتری ثبت‌نام‌شده"
    return "مشتری مهمان"


def _latest_cost_map(product_ids, end_date):
    """Fallback only for historical order items that predate unit_cost."""
    costs = {}
    receipts = (
        Receipt.objects.filter(
            product_id__in=product_ids,
            unit_purchase_price_irr__isnull=False,
            occurred_at__lt=timezone.make_aware(datetime.combine(end_date + timedelta(days=1), time.min)),
        )
        .order_by("product_id", "-occurred_at", "-id")
        .values("product_id", "unit_purchase_price_irr")
    )
    for receipt in receipts:
        costs.setdefault(receipt["product_id"], receipt["unit_purchase_price_irr"])
    return costs


def _row_template(label):
    return {"name": label, "orders": set(), "units": 0, "revenue": ZERO, "cost": ZERO, "profit": ZERO, "known_cost_revenue": ZERO, "estimated_cost_revenue": ZERO}


def _serialize_groups(groups):
    rows = []
    for row in groups.values():
        revenue = row["revenue"]
        coverage = ((row["known_cost_revenue"] + row["estimated_cost_revenue"]) / revenue * Decimal("100")) if revenue else ZERO
        rows.append({
            "name": row["name"], "orders": len(row["orders"]), "units": row["units"],
            "revenue": row["revenue"], "cost": row["cost"], "profit": row["profit"],
            "margin": (row["profit"] / revenue * Decimal("100")) if revenue else ZERO,
            "cost_coverage": coverage,
            "estimated": bool(row["estimated_cost_revenue"]),
        })
    return sorted(rows, key=lambda row: (row["revenue"], row["profit"]), reverse=True)


def build_management_report(params):
    """Return a single source of truth for HTML and XLSX management reports."""
    start, end, stale_days = parse_report_range(params)
    items = list(
        OrderItem.objects.filter(
            order__status__in=COMPLETED_SALE_STATUSES,
            order__created_at__date__gte=start,
            order__created_at__date__lte=end,
        )
        .select_related("order", "order__customer", "product", "product__brand", "product__category")
        .order_by("order__created_at", "id")
    )
    fallback_costs = _latest_cost_map({item.product_id for item in items if item.product_id and item.unit_cost is None}, end)
    by_brand, by_category, by_customer = {}, {}, {}
    daily = defaultdict(lambda: {"revenue": ZERO, "profit": ZERO, "orders": set()})
    total_revenue = total_cost = total_profit = known_cost_revenue = estimated_cost_revenue = ZERO
    order_ids = set()

    for item in items:
        # Cart-level coupon discounts live on Order, not OrderItem. Allocate
        # them proportionally so brand/category/customer sales reconcile to
        # the actual merchandise revenue after discounts (shipping excluded).
        line_total = item.line_total or ZERO
        order_subtotal = item.order.subtotal or ZERO
        coupon_share = (item.order.discount_amount or ZERO) * line_total / order_subtotal if order_subtotal else ZERO
        revenue = line_total - coupon_share
        quantity = item.quantity or 0
        unit_cost = item.unit_cost
        estimated = False
        if unit_cost is None and item.product_id in fallback_costs:
            unit_cost = fallback_costs[item.product_id]
            estimated = True
        cost = (unit_cost * quantity) if unit_cost is not None else ZERO
        profit = revenue - cost if unit_cost is not None else ZERO
        order_ids.add(item.order_id)
        total_revenue += revenue
        total_cost += cost
        total_profit += profit
        if unit_cost is not None:
            if estimated:
                estimated_cost_revenue += revenue
            else:
                known_cost_revenue += revenue

        brand = item.product.brand.name if item.product and item.product.brand_id else "بدون برند"
        category = item.product.category.name_fa if item.product and item.product.category_id else "بدون دسته‌بندی"
        customer = _customer_name(item.order)
        for bucket, label in ((by_brand, brand), (by_category, category), (by_customer, customer)):
            row = bucket.setdefault(label, _row_template(label))
            row["orders"].add(item.order_id)
            row["units"] += quantity
            row["revenue"] += revenue
            row["cost"] += cost
            row["profit"] += profit
            if unit_cost is not None:
                if estimated:
                    row["estimated_cost_revenue"] += revenue
                else:
                    row["known_cost_revenue"] += revenue

        day = item.order.created_at.astimezone(timezone.get_current_timezone()).date()
        daily[day]["revenue"] += revenue
        daily[day]["profit"] += profit
        daily[day]["orders"].add(item.order_id)

    stale_cutoff = timezone.now() - timedelta(days=stale_days)
    stagnant = list(
        Product.objects.filter(is_active=True)
        .select_related("brand", "category", "inventory")
        .annotate(
            last_sale=Max(
                "order_items__order__created_at",
                filter=Q(order_items__order__status__in=COMPLETED_SALE_STATUSES),
            )
        )
        .filter(Q(last_sale__isnull=True) | Q(last_sale__lt=stale_cutoff))
        .order_by("last_sale", "name")
    )
    stale_rows = [
        {
            "code": product.code, "name": product.name,
            "brand": product.brand.name if product.brand_id else "—",
            "category": product.category.name_fa,
            "last_sale": product.last_sale,
            "available": product.inventory.available_quantity if hasattr(product, "inventory") else None,
        }
        for product in stagnant
    ]
    days = []
    day = start
    while day <= end:
        value = daily[day]
        days.append({"date": day, "revenue": value["revenue"], "profit": value["profit"], "orders": len(value["orders"])})
        day += timedelta(days=1)

    cost_coverage = ((known_cost_revenue + estimated_cost_revenue) / total_revenue * Decimal("100")) if total_revenue else ZERO
    return {
        "start": start, "end": end, "stale_days": stale_days,
        "summary": {
            "orders": len(order_ids), "units": sum(item.quantity for item in items), "revenue": total_revenue,
            "cost": total_cost, "profit": total_profit,
            "margin": (total_profit / total_revenue * Decimal("100")) if total_revenue else ZERO,
            "cost_coverage": cost_coverage, "known_cost_revenue": known_cost_revenue,
            "estimated_cost_revenue": estimated_cost_revenue,
            "uncosted_revenue": total_revenue - known_cost_revenue - estimated_cost_revenue,
        },
        "brands": _serialize_groups(by_brand), "categories": _serialize_groups(by_category),
        "customers": _serialize_groups(by_customer), "days": days,
        "stagnant": stale_rows, "stagnant_count": len(stale_rows),
    }


def build_xlsx_response(report):
    """Build a styled, multi-sheet .xlsx workbook for the selected report range."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="0D5C63")
    accent_fill = PatternFill("solid", fgColor="E6F3F0")
    header_font = Font(color="FFFFFF", bold=True)

    def sheet(name, headers, rows, widths=None):
        ws = workbook.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = "A2"
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for values in rows:
            ws.append(values)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(cell.value, (Decimal, float, int)):
                    cell.number_format = '#,##0.00'
        for index, width in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.auto_filter.ref = ws.dimensions
        return ws

    summary = report["summary"]
    overview = workbook.create_sheet("خلاصه")
    overview.sheet_view.rightToLeft = True
    overview["A1"] = "گزارش مدیریتی فروش و سود"
    overview["A1"].font = Font(size=16, bold=True, color="0D5C63")
    overview["A3"] = "بازه گزارش"
    overview["B3"] = f"{report['start'].isoformat()} تا {report['end'].isoformat()}"
    rows = [
        ("تعداد سفارش", summary["orders"]), ("تعداد قلم", summary["units"]),
        ("فروش کالا (ریال)", summary["revenue"]), ("بهای تمام‌شده پوشش‌داده‌شده (ریال)", summary["cost"]),
        ("سود ناخالص (ریال)", summary["profit"]), ("حاشیه سود", summary["margin"] / Decimal("100")),
        ("پوشش بهای خرید", summary["cost_coverage"] / Decimal("100")),
        ("فروش بدون بهای خرید", summary["uncosted_revenue"]), ("کالای راکد", report["stagnant_count"]),
    ]
    for index, (label, value) in enumerate(rows, start=5):
        overview.cell(index, 1, label).fill = accent_fill
        overview.cell(index, 1).font = Font(bold=True)
        overview.cell(index, 2, value)
        overview.cell(index, 2).number_format = "0.0%" if "حاشیه" in label or "پوشش" in label else '#,##0.00'
    overview.column_dimensions["A"].width = 36
    overview.column_dimensions["B"].width = 28

    def group_rows(values):
        return [[row["name"], row["orders"], row["units"], row["revenue"], row["cost"], row["profit"], row["margin"] / Decimal("100"), "برآوردی" if row["estimated"] else "ثبت‌شده"] for row in values]

    group_headers = ["عنوان", "سفارش", "تعداد", "فروش (ریال)", "بهای خرید (ریال)", "سود (ریال)", "حاشیه سود", "منشأ بهای خرید"]
    for name, data in (("برندها", report["brands"]), ("دسته‌ها", report["categories"]), ("مشتریان", report["customers"])):
        ws = sheet(name, group_headers, group_rows(data), [32, 12, 12, 20, 20, 20, 14, 18])
        for cell in ws["G"][1:]:
            cell.number_format = "0.0%"

    sheet("روند روزانه", ["تاریخ", "تعداد سفارش", "فروش (ریال)", "سود (ریال)"], [[row["date"], row["orders"], row["revenue"], row["profit"]] for row in report["days"]], [16, 15, 22, 22])
    sheet("کالای راکد", ["کد", "محصول", "برند", "دسته", "آخرین فروش", "موجودی قابل فروش"], [[row["code"], row["name"], row["brand"], row["category"], row["last_sale"].astimezone(timezone.get_current_timezone()).replace(tzinfo=None) if row["last_sale"] else "بدون فروش", row["available"]] for row in report["stagnant"]], [20, 45, 22, 30, 22, 20])

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="management-report-{report["start"].isoformat()}-{report["end"].isoformat()}.xlsx"'
    return response
